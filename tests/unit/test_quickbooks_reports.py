# Copyright (c) 2026 James Hawkins. PolyForm Noncommercial License 1.0.0 — see LICENSE.md.
"""QuickBooks Online trial balance, A/R aging, inventory valuation and Journal
exports: recognized, flattened, footed, loaded.

The fixtures in tests/fixtures/quickbooks/kestrel_qbo/ are real Excel exports
from a QuickBooks Online Accountant company ("xx") holding a few invented
Kestrel-style transactions, exported 2026-09-30: account 10100 Checking -
First Prairie, two customers, two inventory items, three invoices (April, May,
June 2026), one bill, two checks, one deposit and one journal entry, the
checking account reconciled at 2026-06-30. journal_created_by.xlsx is the
Journal with "Created on" and "Created by" added through Customize.
"""

import io
from pathlib import Path

import openpyxl
import pytest

from assurance_application.service import WorkbenchService
from assurance_artifacts import xlsx
from assurance_artifacts.vault import ArtifactVault
from assurance_persistence.database import connect, migrate
from assurance_persistence.legacy_import import ensure_tenant
from procedures_ap import quickbooks as qb

REAL = Path(__file__).resolve().parents[2] / "tests" / "fixtures" / "quickbooks" / "kestrel_qbo"
TB = (REAL / "trial_balance.xlsx").read_bytes()
AGING = (REAL / "ar_aging_summary.xlsx").read_bytes()
INVENTORY = (REAL / "inventory_valuation_summary.xlsx").read_bytes()
JOURNAL = (REAL / "journal.xlsx").read_bytes()
JOURNAL_CREATED = (REAL / "journal_created_by.xlsx").read_bytes()
LEDGER = (REAL / "general_ledger.xlsx").read_bytes()
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PREPARER, REVIEWER, PARTNER = "p-prep", "p-rev", "p-partner"


def _flatten(content, recipe_id):
    extraction, headers, rows = xlsx.extract(content)
    return qb.apply(recipe_id, headers, rows, extraction["header_row"] + 1)


def _edited(content, edits):
    """A copy of a real export with cells changed (formulas replaced by the
    numbers they held, since a rewritten workbook keeps no cached values)."""
    values = openpyxl.load_workbook(io.BytesIO(content), data_only=True).active
    book = openpyxl.load_workbook(io.BytesIO(content))
    sheet = book.active
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.value = values[cell.coordinate].value
    for ref, value in edits.items():
        sheet[ref] = value
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_each_real_export_is_claimed_by_exactly_its_recipe():
    expected = {TB: ["qbo.trial_balance.trial_balance"],
                AGING: ["qbo.ar_aging_summary.ar_listing"],
                INVENTORY: ["qbo.inventory_valuation_summary.inventory_listing"],
                JOURNAL: ["qbo.journal.journal_entries"],
                JOURNAL_CREATED: ["qbo.journal_created.journal_entries"],
                LEDGER: []}
    for content, recipes in expected.items():
        rows = xlsx.preview(content)[0]["rows"]
        assert [m["recipe"] for m in qb.recognize(rows)] == recipes
        assert all(m["header_row"] == 5 for m in qb.recognize(rows))


def test_the_trial_balance_gives_account_numbers_and_signed_balances():
    headers, rows, source_rows, report = _flatten(TB, "qbo.trial_balance.trial_balance")
    assert headers == ["Account Name", "Debit", "Credit", "Account", "Balance"]
    assert len(rows) == 9 and source_rows == list(range(6, 15))
    assert [(t["column"], t["computed"], t["agrees"]) for t in report["grand_total"]] == \
        [("Debit", "41290.00", True), ("Credit", "41290.00", True)]
    by = {r["Account"]: r["Balance"] for r in rows}
    # An account named with a number is keyed by it; one without, by its name.
    assert by["10100"] == "28915.00"
    assert by["Accounts Payable (A/P)"] == "-1250.00"
    assert by["Opening Balance Equity"] == "-31040.00"
    assert sum(float(v) for v in by.values()) == 0


def test_a_trial_balance_total_that_does_not_foot_is_reported():
    # Cash raised by 100 while the TOTAL keeps its printed figure.
    _, _, _, report = _flatten(_edited(TB, {"B6": 29015}), "qbo.trial_balance.trial_balance")
    [bad] = report["totals_disagreeing"]
    assert (bad["column"], bad["computed"], bad["stated"]) == ("Debit", "41390.00", "41290")


def test_the_aging_keeps_its_buckets_and_foots_every_row():
    headers, rows, _, report = _flatten(AGING, "qbo.ar_aging_summary.ar_listing")
    assert headers[0] == "Customer"
    assert [(r["Customer"], r["CURRENT"], r["1 - 30"], r["31 - 60"], r["Total"])
            for r in rows] == [("Big Sky Pedal Co.", "", "1140", "", "1140"),
                               ("Gallatin Gear Exchange", "1460", "", "1400", "2860")]
    assert report["rows_footed"] == 2 and report["totals_disagreeing"] == []
    assert [t["computed"] for t in report["grand_total"]] == \
        ["1460.00", "1140.00", "1400.00", "0.00", "0.00", "4000.00"]
    # A customer whose Total is not its buckets' sum is reported, not absorbed.
    _, _, _, report = _flatten(_edited(AGING, {"G7": 2960}), "qbo.ar_aging_summary.ar_listing")
    assert {(t["group"], t["column"]) for t in report["totals_disagreeing"]} == \
        {("Gallatin Gear Exchange", "Total"), ("TOTAL", "Total")}


def test_inventory_rows_carry_average_cost_and_nested_categories_are_refused():
    headers, rows, _, report = _flatten(
        INVENTORY, "qbo.inventory_valuation_summary.inventory_listing")
    assert headers == ["Item", "SKU", "Qty", "Asset Value", "Calc. Avg"]
    assert [(r["SKU"], r["Qty"], r["Asset Value"], r["Calc. Avg"]) for r in rows] == \
        [("TIRE-G40", "70", "1540", "22"), ("HUB-TR", "25", "2125", "85")]
    assert [(t["column"], t["agrees"]) for t in report["grand_total"]] == \
        [("Qty", True), ("Asset Value", True)]
    # The layout with categories was not seen in a real export: refused, not guessed.
    nested = _edited(INVENTORY, {"A6": "Tires", "B6": None, "C6": None, "D6": None,
                                 "E6": None})
    with pytest.raises(qb.RecipeError, match="nested heading"):
        _flatten(nested, "qbo.inventory_valuation_summary.inventory_listing")


def test_the_journal_names_entries_numbers_lines_and_checks_each_transaction():
    headers, rows, _, report = _flatten(JOURNAL_CREATED, "qbo.journal_created.journal_entries")
    assert headers[-4:] == ["Entry", "Line", "Account", "Created date"]
    assert report["subtotals_checked"] == 22          # 11 transactions x debit, credit
    assert report["totals_disagreeing"] == []
    entries = list(dict.fromkeys(r["Entry"] for r in rows))
    assert len(entries) == 11 and len(rows) == 31
    assert entries[0] == "2026-03-31 Deposit (no num)"
    assert "2026-04-15 Invoice 1001 Gallatin Gear Exchange" in entries
    assert "2026-05-10 Bill MCC-5521 Moraine Cycle Components" in entries
    je = [r for r in rows if r["Entry"] == "2026-06-30 Journal Entry 1"]
    assert [(r["Line"], r["Account"], r["Debit"], r["Credit"]) for r in je] == \
        [("1", "Bank Charges & Fees", "35", ""), ("2", "10100", "", "35")]
    assert {r["Created date"] for r in rows} == {"09/30/2026"}
    assert {r["Created by"] for r in rows} == {"James Hawkins"}


def test_two_transactions_with_the_same_name_are_told_apart_by_their_id():
    # Both inventory openings are Inventory Starting Value, START, no name,
    # 03/31/2026: QuickBooks' own transaction IDs (3 and 5) keep them apart.
    _, rows, _, _ = _flatten(JOURNAL, "qbo.journal.journal_entries")
    starts = sorted({r["Entry"] for r in rows if "Inventory Starting Value" in r["Entry"]})
    assert starts == ["2026-03-31 Inventory Starting Value START [txn 3]",
                      "2026-03-31 Inventory Starting Value START [txn 5]"]


def test_a_journal_transaction_that_does_not_balance_is_reported():
    # Invoice 1001's A/R line raised by 1: its 'Total for 4' no longer agrees.
    _, _, _, report = _flatten(_edited(JOURNAL, {"H19": 1401}), "qbo.journal.journal_entries")
    assert {(t["group"], t["column"]) for t in report["totals_disagreeing"]} == \
        {("4", "Debit"), ("TOTAL", "Debit")}


def test_real_period_lines_name_their_end_date():
    # QuickBooks' Excel titles abbreviate months, and a ledger or journal run
    # for whole months names no day: both used to read as undated.
    assert qb.last_date("As of Jun 30, 2026") == "2026-06-30"
    assert qb.last_date("April-June, 2026") == "2026-06-30"
    assert qb.last_date("January-June, 2026") == "2026-06-30"
    assert qb.last_date("August 2026") == "2026-08-31"
    assert qb.last_date("February 2028") == "2028-02-29"
    assert qb.last_date("January 1-September 23, 2026") == "2026-09-23"
    ledger = qb.gl_account_balance(xlsx.preview(LEDGER, max_rows=500)[0]["rows"])
    assert (ledger["as_of"], ledger["beginning"], ledger["ending"]) == \
        ("2026-06-30", "0.00", "1250.00")


@pytest.fixture()
def service(tmp_path):
    conn = connect(tmp_path / "control.db")
    migrate(conn)
    svc = WorkbenchService(conn, ArtifactVault(tmp_path / "vault"),
                           ensure_tenant(conn, "qbo-tb"))
    eid = svc.create_engagement(PARTNER, "xx", "2026-06-30")["engagement_id"]
    svc.assign_team(PARTNER, eid, PREPARER, "preparer")
    svc.assign_team(PARTNER, eid, REVIEWER, "reviewer")
    yield svc, eid
    conn.close()


def _store(svc, eid, content, name):
    return svc.store_source(PREPARER, eid, content=content, media_type=XLSX_TYPE,
                            original_name=name)["artifact_id"]


def test_the_trial_balance_is_built_from_this_year_and_last_and_loaded(service):
    svc, eid = service
    prior = _edited(TB, {"A3": "As of Jun 30, 2025", "B6": 20000, "C10": 22125,
                         "B15": 32375, "C15": 32375})
    current_id = _store(svc, eid, TB, "Trial Balance 2026.xlsx")
    prior_id = _store(svc, eid, prior, "Trial Balance 2025.xlsx")
    found = svc.trial_balance_candidates(eid)["trial_balances"]
    assert {(c["artifact_id"], c["as_of"]) for c in found} == \
        {(current_id, "2026-06-30"), (prior_id, "2025-06-30")}

    built = svc.build_trial_balance(PREPARER, eid, current_artifact_id=current_id,
                                    prior_artifact_id=prior_id)
    assert (built["as_of"], built["prior_as_of"], built["accounts"]) == \
        ("2026-06-30", "2025-06-30", 9)
    assert built["notes"] == []
    proposal = svc.propose_source_mapping(PREPARER, eid, role="Trial_balance",
                                          artifact_id=built["artifact_id"])
    assert proposal["column_map"] == {"account": "Account", "description": "Description",
                                      "balance": "Balance", "prior_balance": "Prior Balance"}
    svc.approve_source_mapping(REVIEWER, eid, proposal["spec_id"])
    svc.normalize_source(PREPARER, eid, proposal["spec_id"])
    records = svc._rebuild_table(proposal["spec_id"]).records
    cash = next(r for r in records if r["account"] == "10100")
    assert (str(cash["balance"]), str(cash["prior_balance"])) == ("28915.00", "20000.00")

    # The prior must be the earlier one; the wrong way round is refused.
    with pytest.raises(ValueError, match="must be dated before"):
        svc.build_trial_balance(PREPARER, eid, current_artifact_id=prior_id,
                                prior_artifact_id=current_id)
    # One that does not foot is refused before anything is stored.
    broken = _store(svc, eid, _edited(TB, {"B6": 1}), "Trial Balance broken.xlsx")
    with pytest.raises(ValueError, match="does not foot"):
        svc.build_trial_balance(PREPARER, eid, current_artifact_id=broken)


def test_the_journal_and_aging_load_through_their_recipes(service):
    svc, eid = service
    for content, name, role, recipe in (
            (JOURNAL_CREATED, "Journal.xlsx", "Journal_entries",
             "qbo.journal_created.journal_entries"),
            (AGING, "AR Aging Summary.xlsx", "AR_listing", "qbo.ar_aging_summary.ar_listing"),
            (INVENTORY, "Inventory Valuation Summary.xlsx", "Inventory_listing",
             "qbo.inventory_valuation_summary.inventory_listing")):
        artifact_id = _store(svc, eid, content, name)
        proposal = svc.propose_source_mapping(PREPARER, eid, role=role,
                                              artifact_id=artifact_id,
                                              extraction={"recipe": recipe})
        svc.approve_source_mapping(REVIEWER, eid, proposal["spec_id"])
        rec = svc.normalize_source(PREPARER, eid, proposal["spec_id"])["reconciliation"]
        assert rec["rows_rejected"] == 0, (name, rec)
    journal = next(d for d in svc.sources(eid)["datasets"] if d["role"] == "Journal_entries")
    assert journal["rows_loaded"] == 31


def test_an_added_journal_that_repeats_loaded_entries_is_refused(service):
    # Review M3: a re-exported Journal (new footer, so new bytes) added to the
    # one in use would merge each entry's lines and double-count them.
    svc, eid = service

    def load(content, name, mode=None):
        proposal = svc.propose_source_mapping(
            PREPARER, eid, role="Journal_entries", artifact_id=_store(svc, eid, content, name),
            extraction={"recipe": "qbo.journal_created.journal_entries"})
        svc.approve_source_mapping(REVIEWER, eid, proposal["spec_id"])
        return svc.normalize_source(PREPARER, eid, proposal["spec_id"], mode=mode)

    load(JOURNAL_CREATED, "Journal.xlsx")
    again = _edited(JOURNAL_CREATED, {"A63": " Thursday, October 1, 2026 08:00 AM GMT-07:00"})
    with pytest.raises(ValueError, match=r"11 of its entries are already loaded: .*\(in Journal.xlsx\)"):
        load(again, "Journal re-export.xlsx", mode="add")
    journal = [d for d in svc.sources(eid)["datasets"] if d["role"] == "Journal_entries"]
    assert [d["rows_loaded"] for d in journal] == [31]


def test_a_line_mapping_counts_only_when_it_reaches_a_loaded_account(service):
    # Review M1: a mapping for an account the trial balance does not hold
    # supplies no line, so coverage must not call the analytics executable.
    svc, eid = service
    svc.update_workflow(PARTNER, eid, "cycles", {"cycles": ["planning"]})
    built = svc.build_trial_balance(PREPARER, eid,
                                    current_artifact_id=_store(svc, eid, TB, "TB.xlsx"))
    proposal = svc.propose_source_mapping(PREPARER, eid, role="Trial_balance",
                                          artifact_id=built["artifact_id"])
    svc.approve_source_mapping(REVIEWER, eid, proposal["spec_id"])
    svc.normalize_source(PREPARER, eid, proposal["spec_id"])

    def status():
        rows = svc.coverage(eid)["procedures"]
        return next(r for r in rows if r["procedure_id"] == "fs.trial_balance_analytics")["status"]
    assert status() == "partial"
    svc.update_workflow(PREPARER, eid, "line_mapping", {"account": "99999", "line": "cash"})
    assert status() == "partial"
    svc.update_workflow(PREPARER, eid, "line_mapping", {"account": "10100", "line": "cash"})
    assert status() == "executable"
