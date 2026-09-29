# Copyright (c) 2026 James Hawkins. PolyForm Noncommercial License 1.0.0 — see LICENSE.md.
"""Kestrel Valley part 1: write the QuickBooks payables exports, the Journal
report and the part-1 answer key from generate_payables.build().

    python case-studies/kestrel-valley-cycle/payables_outputs.py
"""

from __future__ import annotations

import json
from datetime import timedelta
from decimal import Decimal
from decimal import Decimal as D

import openpyxl

import generate as g
from generate_payables import (
    ACCTS, BOOKKEEPER, CUR, FY_START, HOLIDAYS, PE, PRI, VENDORS, ZERO, bal, build, m,
)

NAMES = {n: f"{n} {name}" for n, name, *_ in g.TB}
AP_NAME, CHECKING = NAMES["20000"], NAMES["10100"]
STAMP = "Wednesday, July 15, 2026 10:50 AM GMTZ"
VENDOR_NAMES = {v[0] for v in VENDORS}


def write_vendor_list():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in ([g.COMPANY], ["Vendor Contact List"], [None],
                ["Vendor", "Phone numbers", "Email", "Full name", "Billing address",
                 "Account #"], *[list(v) for v in sorted(VENDORS)],
                [None], [None], [None], [" " + STAMP]):
        ws.append(row)
    g.save(wb, g.QBO / "Vendor_Contact_List.xlsx")


def write_transaction_list(b):
    rows_by_vendor: dict[str, list] = {}
    for vendor, when, po, amount in b.pos:
        rows_by_vendor.setdefault(vendor, []).append(
            (when, "No", "Purchase Order", po, "No", "", AP_NAME, "", amount))
    for t in b.tx:
        if t["type"] not in ("Bill", "Bill Payment (Check)", "Check"):
            continue
        if t["name"] not in VENDOR_NAMES:
            continue
        if t["type"] == "Bill Payment (Check)":
            other = "20000"
        else:
            other = next(n for n, a in t["lines"] if n not in ("20000", "10100"))
        amount = sum(a for _, a in t["lines"] if a > 0)
        track = "Yes" if t["name"] in ("Alder & Finch CPAs", "DM Consulting",
                                       "Hyalite Fabrication") else "No"
        if t["type"] == "Bill":
            row = (t["date"], track, "Bill", t["num"], "Yes", t["memo"], AP_NAME,
                   NAMES[other], amount)
        else:
            row = (t["date"], track, t["type"], t["num"], "Yes", t["memo"], CHECKING,
                   NAMES[other], -amount)
        rows_by_vendor.setdefault(t["name"], []).append(row)
    body, grand = [], ZERO
    for vendor in sorted(rows_by_vendor):
        rows = sorted(rows_by_vendor[vendor], key=lambda r: (r[0], r[2], str(r[3])))
        body.append((vendor,) + (None,) * 9)
        for r in rows:
            body.append((None, r[0].strftime("%m/%d/%Y"), *r[1:]))
        total = sum(r[8] for r in rows)
        grand += total
        body.append((f"Total for {vendor}",) + (None,) * 8 + (total,))
    body.append(("TOTAL",) + (None,) * 8 + (grand,))
    g.write_report(g.QBO / "Transaction_List_by_Vendor.xlsx", "Transaction List by Vendor",
                   "July 1, 2025-June 30, 2026",
                   [None, "Date", "Track 1099", "Transaction type", "Num",
                    "Posting (Y/N)", "Memo", "Account full name", "Item split account",
                    "Amount"], body, footer_basis=False, stamp=STAMP)


def write_bill_payment_list(b):
    rows = sorted((t for t in b.tx if t["type"] == "Bill Payment (Check)"),
                  key=lambda t: (t["date"], int(t["num"])))
    body = [(CHECKING, None, None, None, None)]
    total = ZERO
    for t in rows:
        amount = -sum(a for _, a in t["lines"] if a > 0)
        total += amount
        body.append((None, t["date"].strftime("%m/%d/%Y"), t["num"], t["name"], amount))
    body.append((f"Total for {CHECKING}", None, None, None, total))
    g.write_report(g.QBO / "Bill_Payment_List.xlsx", "Bill Payment List",
                   "July 1, 2025-June 30, 2026", [None, "Date", "Num", "Vendor", "Amount"],
                   body, footer_basis=False, stamp=STAMP)


def open_bills(bills, end_ap):
    """Allocate each vendor's year-end balance to its latest bills."""
    out = []
    for vendor, balance in sorted(end_ap.items()):
        left = balance
        for _, when, num, amount, _, _ in sorted(
                (x for x in bills if x[0] == vendor), key=lambda x: x[1], reverse=True):
            if left <= 0:
                break
            open_amount = min(amount, left)
            out.append((vendor, when, num, amount, open_amount))
            left -= open_amount
        assert left == 0, (vendor, left)
    return out


def write_unpaid_bills(open_list):
    body, total_amt, total_open = [], ZERO, ZERO
    for vendor in sorted({o[0] for o in open_list}):
        rows = sorted((o for o in open_list if o[0] == vendor), key=lambda o: o[1])
        body.append((vendor,) + (None,) * 7)
        for _, when, num, amount, open_amount in rows:
            due = when + timedelta(days=30)
            body.append((None, when.strftime("%m/%d/%Y"), "Bill", num,
                         due.strftime("%m/%d/%Y"), float((PE - due).days), amount,
                         open_amount))
        a, o = sum(r[3] for r in rows), sum(r[4] for r in rows)
        total_amt, total_open = total_amt + a, total_open + o
        body.append((f"Total for {vendor}",) + (None,) * 5 + (a, o))
    body.append(("TOTAL",) + (None,) * 5 + (total_amt, total_open))
    g.write_report(g.QBO / "Unpaid_Bills.xlsx", "Unpaid Bills Report", "As of June 30, 2026",
                   [None, "Date", "Transaction type", "Num", "Due date", "Past due",
                    "Amount", "Open balance"], body, footer_basis=False, stamp=STAMP)
    return total_open


def write_journal(b):
    """QuickBooks Online Journal report (modeled): each transaction's first
    line carries its date, type, num and name; a debit/credit subtotal closes
    it; Create date and Created by are the customized columns."""
    body = []
    txs = sorted(b.tx, key=lambda t: (t["date"], t["type"], str(t["num"] or ""),
                                      t["name"]))
    for t in txs:
        for i, (n, amount) in enumerate(t["lines"]):
            debit = amount if amount > 0 else None
            credit = -amount if amount < 0 else None
            if i == 0:
                body.append((None, t["date"].strftime("%m/%d/%Y"), t["type"],
                             t["num"] or "", t["name"], t["memo"], NAMES[n], debit, credit,
                             t["created"].strftime("%m/%d/%Y"), t["by"]))
            else:
                body.append((None, None, None, None, None, t["memo"], NAMES[n], debit,
                             credit, None, None))
        dr = sum(a for _, a in t["lines"] if a > 0)
        body.append((None, None, None, None, None, None, None, dr, dr, None, None))
    total = sum(sum(a for _, a in t["lines"] if a > 0) for t in txs)
    body.append(("TOTAL", None, None, None, None, None, None, total, total, None, None))
    g.write_report(g.QBO / "Journal.xlsx", "Journal", "July 1, 2025-June 30, 2026",
                   [None, "Date", "Transaction type", "Num", "Name", "Memo/Description",
                    "Account", "Debit", "Credit", "Create date", "Created by"], body,
                   stamp=STAMP)
    return txs


def je_selections(txs):
    """The journal-entry tests as the key defines them, in plain logic. An
    entry is one transaction: (date, type, num, name)."""
    use: dict[str, set] = {}
    for i, t in enumerate(txs):
        for n, _ in t["lines"]:
            use.setdefault(n, set()).add(i)
    rare = {n for n, s in use.items() if len(s) <= 1}
    out = {k: [] for k in ("posted_after_period_end", "weekend_or_holiday",
                           "round_amount", "unauthorized_user", "seldom_used_account",
                           "no_description")}
    for t in txs:
        ident = f"{t['date']} {t['type']} {t['num'] or '(no num)'} {t['name']}".strip()
        debits = sum(a for _, a in t["lines"] if a > 0)
        if t["date"] <= PE < t["created"]:
            out["posted_after_period_end"].append(ident)
        if t["created"].weekday() >= 5 or t["created"] in HOLIDAYS:
            out["weekend_or_holiday"].append(ident)
        if debits >= 10000 and debits % 1000 == 0:
            out["round_amount"].append(ident)
        if t["by"] != BOOKKEEPER:
            out["unauthorized_user"].append(ident)
        if any(n in rare for n, _ in t["lines"]):
            out["seldom_used_account"].append(ident)
        if not t["memo"]:
            out["no_description"].append(ident)
    manual_no_memo = [f"{t['type']} {t['num']}" for t in txs
                      if t["type"] == "Journal Entry" and not t["memo"]]
    return ({k: {"count": len(v), "entries": v if len(v) <= 15 else v[:15] + ["..."]}
             for k, v in out.items()}, sorted(rare), manual_no_memo)


def main():
    b, key, act, bills, payments, begin_ap, end_ap = build()
    activity: dict[str, Decimal] = {}
    for t in b.tx:
        for n, amount in t["lines"]:
            activity[n] = activity.get(n, ZERO) + amount
    gaps = {n: activity.get(n, ZERO) - act[n] for n in act
            if n != "32000" and activity.get(n, ZERO) != act[n]}
    assert not gaps, gaps
    june = sum(a for t in b.tx if (t["date"].year, t["date"].month) == (2026, 6)
               for n, a in t["lines"] if n == "10100")
    rec = (sum(m(d[4]) for d in g.CLEARED_DEPOSITS + g.UNCLEARED_DEPOSITS)
           - sum(m(c[4]) for c in g.CLEARED_CHECKS + g.UNCLEARED_CHECKS))
    assert june == rec, (june, rec)

    write_vendor_list()
    write_transaction_list(b)
    write_bill_payment_list(b)
    unpaid = write_unpaid_bills(open_bills(bills, end_ap))
    assert unpaid == -bal(CUR, "20000"), unpaid
    txs = write_journal(b)
    selections, rare, manual_no_memo = je_selections(txs)
    closing = sum(bal(PRI, n) for n in ACCTS if n[0] in "456")

    key.update({
        "case_part": "1 of 3: payables and the journal",
        "period": {"start": FY_START.isoformat(), "end": PE.isoformat()},
        "vendors": {"count": len(VENDORS),
                    "twins": [["Moraine Cycle Components", "Moraine Cycle Components, Inc.",
                               "same address and phone"]]},
        "bills": {"count": len(bills), "total": str(sum(x[3] for x in bills))},
        "bill_payments": {"count": len(payments),
                          "total": str(sum(p[3] for p in payments))},
        "ap_subledger_to_ledger": {"unpaid_bills_total": str(unpaid),
                                   "tb_accounts_payable": str(-bal(CUR, "20000")),
                                   "difference": "0.00"},
        "short_payment": {"vendor": "Alder & Finch CPAs", "bill": "1870.00",
                          "check_4425": "1780.00", "left_open": "90.00"},
        "purchase_orders": {"count": len(b.pos),
                            "three_way_match": "not testable: QuickBooks Online "
                                               "records no receipts"},
        "segregation_of_duties": "not testable: the exports carry no approver",
        "journal": {"transactions": len(txs),
                    "lines": sum(len(t["lines"]) for t in txs),
                    "entry_identity": "(date, transaction type, num, name): Num alone "
                                      "is not unique and is blank for deposits, "
                                      "transfers and payroll"},
        "je_testing_policies": {"je_authorized_users": BOOKKEEPER,
                                "je_round_amount_threshold": "10000",
                                "je_round_unit": "1000", "je_seldom_used_max": "1",
                                "je_holidays": ", ".join(sorted(str(h)
                                                                for h in HOLIDAYS))},
        "je_testing": selections,
        "je_seldom_used_accounts": rare,
        "je_manual_entries_without_description": manual_no_memo,
        "je_population_completeness": {
            "every_account_rolls_forward": True, "closed_to": "32000",
            "prior_year_closing_amount": str(closing)},
        "misstatement_from_duplicate": {
            "amount": key["duplicate_bill"]["duplicate_amount"],
            "effect": "purchases and cost of sales overstated; a refund is due"},
    })
    # A second path: the key's totals re-read from the files just written.
    ws = openpyxl.load_workbook(g.QBO / "Transaction_List_by_Vendor.xlsx").active
    in_file = {"Bill": ZERO, "Bill Payment (Check)": ZERO}
    for r in ws.iter_rows(min_row=6, values_only=True):
        if r[3] in in_file:
            in_file[r[3]] += abs(m(r[9]))
    assert in_file["Bill"] == D(key["bills"]["total"]), in_file
    assert in_file["Bill Payment (Check)"] == D(key["bill_payments"]["total"]), in_file
    (g.INSTRUCTOR / "answer_key_payables.json").write_text(
        json.dumps(key, indent=2, default=str) + "\n", encoding="utf-8")
    print(json.dumps({k: key[k] for k in ("bills", "bill_payments",
                                          "ap_subledger_to_ledger", "journal")}, indent=1))
    print({k: v["count"] for k, v in selections.items()}, rare, manual_no_memo)


if __name__ == "__main__":
    main()
